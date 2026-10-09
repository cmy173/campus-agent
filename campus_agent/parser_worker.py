"""Executed in a separate process with a wall-time limit. No OCR/network/tool execution."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
import json
import sys

MAX_UNITS = 200
MAX_TEXT = 600_000

def parse(content: bytes, extension: str):
    units, warnings = [], []
    def add(label, text):
        cleaned = text.replace('\x00', '').strip()
        if cleaned:
            units.append({'locator': label, 'text': cleaned})
        if sum(len(x['text']) for x in units) > MAX_TEXT:
            raise ValueError('提取文字超过 60 万字限制，请拆分文件')

    if extension == '.txt':
        try:
            text = content.decode('utf-8-sig')
        except UnicodeDecodeError:
            raise ValueError('TXT 请使用 UTF-8 编码')
        if '\x00' in text:
            raise ValueError('TXT 含二进制内容')
        for n, line in enumerate(text.splitlines(), 1):
            add(f'第 {n} 行', line)
    elif extension == '.pdf':
        if not content.startswith(b'%PDF-'):
            raise ValueError('文件内容不是 PDF')
        from pypdf import PdfReader
        reader = PdfReader(BytesIO(content))
        if reader.is_encrypted:
            raise ValueError('暂不支持加密 PDF')
        if len(reader.pages) > MAX_UNITS:
            raise ValueError('PDF 超过 200 页限制')
        for n, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            if not text.strip():
                warnings.append(f'第 {n} 页没有可提取文字；图片和扫描件暂不进行 OCR')
            add(f'第 {n} 页', text)
    elif extension in ('.docx', '.pptx'):
        try:
            with ZipFile(BytesIO(content)) as z:
                info = z.infolist()
                if len(info) > 3000 or sum(x.file_size for x in info) > 50 * 1024 * 1024:
                    raise ValueError('Office 文件解压规模超限')
                if any('vbaproject' in x.filename.lower() for x in info):
                    raise ValueError('不接受带宏的文件')
                required = 'word/document.xml' if extension == '.docx' else 'ppt/presentation.xml'
                if required not in z.namelist():
                    raise ValueError('文件内容与扩展名不符')
        except BadZipFile:
            raise ValueError('Office 文件损坏或格式不正确')
        if extension == '.docx':
            from docx import Document
            doc = Document(BytesIO(content))
            for n, paragraph in enumerate(doc.paragraphs, 1):
                add(f'第 {n} 段', paragraph.text)
            for n, table in enumerate(doc.tables, 1):
                for row, item in enumerate(table.rows, 1):
                    add(f'表 {n} 第 {row} 行', ' | '.join(c.text for c in item.cells))
            warnings.append('DOCX 按段落及表格定位，不提供排版页码；页眉页脚、文本框和批注不在提取范围')
        else:
            from pptx import Presentation
            deck = Presentation(BytesIO(content))
            if len(deck.slides) > MAX_UNITS:
                raise ValueError('PPTX 超过 200 页限制')
            for n, slide in enumerate(deck.slides, 1):
                parts = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        parts.append(shape.text)
                    if shape.has_table:
                        parts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
                add(f'幻灯片 {n}', '\n'.join(parts))
            warnings.append('PPTX 仅提取文本框和表格；图片、组合图形、备注和音视频不在提取范围')
    else:
        raise ValueError('仅支持 TXT、PDF、DOCX 和 PPTX')
    if not units:
        raise ValueError('没有可提取的文字；扫描件、纯图片或空文件暂不支持')
    # Chunks retain exact offsets and never invent page locations.
    chunks = []
    for unit in units:
        text = unit['text']
        for start in range(0, len(text), 1100):
            chunks.append({'locator': unit['locator'] + (f' · 字符 {start+1}–{min(start+1100,len(text))}' if len(text) > 1100 else ''),
                           'text': text[start:start+1100]})
    return {'chunks': chunks, 'warnings': warnings}

if __name__ == '__main__':
    try:
        result = parse(sys.stdin.buffer.read(), sys.argv[1])
        sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False).encode('utf-8'))
    except ValueError as exc:
        sys.stdout.buffer.write(json.dumps({'error': str(exc)}, ensure_ascii=False).encode('utf-8'))
    except Exception:
        sys.stdout.buffer.write(json.dumps({'error': '解析失败：文件损坏、格式不兼容或缺少解析依赖'}, ensure_ascii=False).encode('utf-8'))

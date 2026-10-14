"""Portable document renderer invoked by the product's ordinary Python tool.

No model calls and no Codex dependencies. It renders supplied facts without
inventing dates or summarizing sources. Existing output files are preserved.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
import re
import sys
import tempfile
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape


def text(value):
    if isinstance(value, dict):
        return value.get('date', value.get('formula', ''))
    return '' if value is None else str(value)


def validate(spec, formats):
    if not isinstance(spec, dict) or not isinstance(spec.get('title'), str) or not spec['title'].strip():
        raise ValueError('title 必须是非空字符串')
    if set(formats) - {'pdf', 'xlsx', 'md'} or not formats:
        raise ValueError('formats 仅支持 pdf,xlsx,md')
    if {'pdf', 'md'} & set(formats) and not spec.get('sections'):
        raise ValueError('PDF/Markdown 至少需要一个 section')
    if 'xlsx' in formats and not spec.get('sheets'):
        raise ValueError('Excel 至少需要一张 sheet')
    for section in spec.get('sections', []):
        if not isinstance(section, dict):
            raise ValueError('section 必须是对象')
        for key in ('paragraphs', 'bullets'):
            if key in section and (not isinstance(section[key], list) or any(not isinstance(v, str) for v in section[key])):
                raise ValueError(f'{key} 必须是字符串数组')
    tables = [section['table'] for section in spec.get('sections', []) if 'table' in section] + spec.get('sheets', [])
    names = set()
    for sheet in spec.get('sheets', []):
        name = sheet.get('name', '')
        if not isinstance(name, str) or not name.strip() or len(name) > 31 or re.search(r'[\\/*?:\[\]]', name) or name.casefold() in names:
            raise ValueError('工作表名称必须唯一，长度1–31且不能含 Excel 禁用字符')
        names.add(name.casefold())
    for table in tables:
        columns, rows = table.get('columns'), table.get('rows')
        if not isinstance(columns, list) or not columns or any(not isinstance(v, str) or not v for v in columns):
            raise ValueError('columns 必须是非空标题数组')
        if not isinstance(rows, list) or any(not isinstance(r, list) or len(r) != len(columns) for r in rows):
            raise ValueError('每行单元格数量必须与 columns 一致')
        if len(rows) > 1_048_570 or len(columns) > 16_384:
            raise ValueError('数据超出 Excel 工作表容量')
        for value in columns + [v for row in rows for v in row]:
            if isinstance(value, dict):
                if set(value) == {'date'}:
                    date.fromisoformat(value['date'])
                elif set(value) == {'formula'} and isinstance(value['formula'], str) and value['formula'].startswith('='):
                    pass
                else:
                    raise ValueError('单元格对象仅支持明确的 date 或 formula')
            elif not isinstance(value, (str, int, float, bool, type(None))):
                raise ValueError('不支持的单元格类型')
            if isinstance(value, float) and not math.isfinite(value):
                raise ValueError('数值必须有限，缺失数据请使用 null')
            if len(text(value)) > 32767:
                raise ValueError('单元格超过32767字符，请拆分内容')


def markdown(spec, path):
    lines = ['# ' + spec['title'], '', spec.get('subtitle', ''), '']
    for section in spec['sections']:
        if section.get('heading'):
            lines += ['## ' + section['heading'], '']
        for para in section.get('paragraphs', []):
            lines += [para, '']
        lines += ['- ' + item for item in section.get('bullets', [])] + ['']
        if table := section.get('table'):
            def row(values):
                return '| ' + ' | '.join(text(v).replace('|', '\\|').replace('\n', '<br>') for v in values) + ' |'
            lines += [row(table['columns']), row(['---'] * len(table['columns']))]
            lines += [row(r) for r in table['rows']] + ['']
    if spec.get('sources'):
        lines += ['## 来源', ''] + ['- ' + source for source in spec['sources']]
    path.write_text('\n'.join(lines).rstrip() + '\n', encoding='utf-8')
    return {}


def pdf(spec, path, font_path=None):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    from pypdf import PdfReader

    candidates = [Path(font_path)] if font_path else [Path(os.environ.get('SystemRoot', 'C:/Windows')) / 'Fonts/msyh.ttc',
        Path('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')]
    font_file = next((p for p in candidates if p.is_file()), None)
    if font_file is None:
        raise ValueError('未找到可嵌入字体，请通过 --font 指定 TrueType 字体')
    pdfmetrics.registerFont(TTFont('CampusDocument', str(font_file), subfontIndex=0))
    chars = set(''.join(text(v) for v in [spec['title'], spec.get('subtitle', '')]))
    for section in spec['sections']:
        chars.update(str(section))
    chars.update(str(spec.get('sources', [])))
    face = pdfmetrics.getFont('CampusDocument').face
    missing = sorted(c for c in chars if not c.isspace() and ord(c) not in face.charToGlyph)
    if missing:
        raise ValueError('字体缺少字符，请选择合适字体：' + ''.join(missing[:20]))
    wide = any(len(section.get('table', {}).get('columns', [])) > 5 for section in spec['sections'])
    page_size = landscape(A4) if wide else A4
    doc = SimpleDocTemplate(str(path), pagesize=page_size, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=44,
                            title=spec['title'], author='Campus Agent')
    base = ParagraphStyle('body', fontName='CampusDocument', fontSize=10.5, leading=16, spaceAfter=9,
                          wordWrap='CJK', alignment=TA_LEFT, textColor=colors.HexColor('#243247'))
    title = ParagraphStyle('title', parent=base, fontSize=20, leading=28, spaceAfter=14)
    heading = ParagraphStyle('heading', parent=base, fontSize=13, leading=20, spaceBefore=12, spaceAfter=8, keepWithNext=True)
    note = ParagraphStyle('note', parent=base, fontSize=9, leading=14, textColor=colors.HexColor('#5D6C80'))
    cell = ParagraphStyle('cell', parent=base, fontSize=9.5, leading=14, spaceAfter=0)
    header = ParagraphStyle('header', parent=cell, textColor=colors.white)
    def para(value, style=base):
        return Paragraph(escape(text(value)).replace('\n', '<br/>'), style)
    flow = [para(spec['title'], title)]
    if spec.get('subtitle'):
        flow.append(para(spec['subtitle'], note))
    for section in spec['sections']:
        if section.get('heading'):
            flow.append(para(section['heading'], heading))
        flow.extend(para(p) for p in section.get('paragraphs', []))
        flow.extend(para('• ' + p) for p in section.get('bullets', []))
        if table := section.get('table'):
            data = [[para(v, header) for v in table['columns']]] + [[para(v, cell) for v in row] for row in table['rows']]
            result = LongTable(data, colWidths=[doc.width / len(table['columns'])] * len(table['columns']), repeatRows=1,
                               splitByRow=1, splitInRow=1, hAlign='LEFT')
            result.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#243C5A')),
                ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F6FA')]),
                ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8), ('TOPPADDING', (0, 0), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 8)]))
            flow += [result, Spacer(1, 10)]
    if spec.get('sources'):
        flow += [para('来源', heading)] + [para(s, note) for s in spec['sources']]
    def footer(canvas, _doc):
        canvas.saveState()
        canvas.setFont('CampusDocument', 8)
        canvas.setFillColor(colors.HexColor('#65758A'))
        canvas.drawRightString(page_size[0] - 42, 24, str(_doc.page))
        canvas.restoreState()
    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    with path.open('rb') as stream:
        reader = PdfReader(stream)
        if not reader.pages or not ''.join(p.extract_text() or '' for p in reader.pages).strip():
            raise ValueError('生成的 PDF 没有可读取的正文')
        return {'pages': len(reader.pages)}


def xlsx(spec, path):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    book.calculation.fullCalcOnLoad = True
    book.calculation.forceFullCalc = True
    book.calculation.calcMode = 'auto'
    counts = {}
    def set_cell(cell, value):
        if isinstance(value, dict):
            if 'date' in value:
                cell.value = date.fromisoformat(value['date'])
                cell.number_format = 'yyyy-mm-dd'
            else:
                cell.value = value['formula']
        else:
            cell.value = value
            if isinstance(value, str):
                cell.data_type = 's'
        cell.font = Font(name='Microsoft YaHei', size=10, color='243247')
        numeric = isinstance(value, (int, float)) and not isinstance(value, bool) or isinstance(value, dict)
        cell.alignment = Alignment(horizontal='right' if numeric else 'left', vertical='top', wrap_text=True, indent=1)
    for item in spec['sheets']:
        sheet = book.create_sheet(item['name'])
        sheet.sheet_view.showGridLines = False
        set_cell(sheet.cell(2, 1), spec['title'])
        sheet.cell(2, 1).font = Font(name='Microsoft YaHei', size=15, bold=True, color='243C5A')
        sheet.row_dimensions[2].height = 26
        for row_index, values in enumerate([item['columns']] + item['rows'], 4):
            max_lines = 1
            for col, value in enumerate(values, 1):
                cell = sheet.cell(row_index, col)
                set_cell(cell, value)
                length = sum(2 if ord(c) > 127 else 1 for c in text(value))
                max_lines = max(max_lines, math.ceil(length / 26), text(value).count('\n') + 1)
                if row_index == 4:
                    cell.fill = PatternFill('solid', fgColor='243C5A')
                    cell.font = Font(name='Microsoft YaHei', size=10, bold=True, color='FFFFFF')
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                    cell.border = Border(right=Side(style='thin', color='FFFFFF'))
                else:
                    cell.border = Border(right=Side(style='thin', color='DDE4ED'))
                    if row_index % 2:
                        cell.fill = PatternFill('solid', fgColor='F3F6FA')
            sheet.row_dimensions[row_index].height = max(26, max_lines * 16 + 8)
        for col in range(1, len(item['columns']) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 28
            sheet.cell(2, col).border = Border(bottom=Side(style='thin', color='BCC8D5'))
        if len(item['rows']) > 15:
            sheet.freeze_panes = 'A5'
        last = len(item['rows']) + 4
        sheet.auto_filter.ref = f'A4:{get_column_letter(len(item["columns"]))}{last}'
        for offset, source in enumerate(item.get('sources', spec.get('sources', [])), last + 3):
            set_cell(sheet.cell(offset, 1), '来源：' + source)
            sheet.row_dimensions[offset].height = max(26, math.ceil(len(source) / 13) * 16 + 8)
        counts[item['name']] = len(item['rows'])
    book.save(path)
    checked = load_workbook(path, read_only=True, data_only=False)
    try:
        if checked.sheetnames != list(counts):
            raise ValueError('工作表导出不完整')
        for name, count in counts.items():
            if checked[name].max_row < count + 4:
                raise ValueError('工作表数据行缺失')
    finally:
        checked.close()
    return {'sheets': counts}


def export(spec, output_dir, stem, formats=('pdf', 'xlsx', 'md'), font_path=None):
    formats = list(dict.fromkeys(formats))
    validate(spec, formats)
    if not stem or stem in {'.', '..'} or re.search(r'[<>:"/\\|?*\x00-\x1f]', stem) or stem.endswith(('.', ' ')):
        raise ValueError('stem 必须是有效文件名，不包含目录')
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    targets = {fmt: output_dir / f'{stem}.{fmt}' for fmt in formats}
    if any(path.exists() for path in targets.values()):
        raise FileExistsError('已有同名成果，请使用新名称保留旧版本')
    results = []
    # Build in an owned sibling directory. Publish only after every renderer succeeds.
    with tempfile.TemporaryDirectory(prefix='.document-build-', dir=output_dir) as scratch:
        for fmt in formats:
            path = Path(scratch) / targets[fmt].name
            details = pdf(spec, path, font_path) if fmt == 'pdf' else xlsx(spec, path) if fmt == 'xlsx' else markdown(spec, path)
            results.append({'path': str(targets[fmt]), 'size': path.stat().st_size,
                            'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), **details})
        for fmt in formats:
            # Exclusive creation also preserves an output created while rendering.
            with targets[fmt].open('xb') as output, (Path(scratch) / targets[fmt].name).open('rb') as source:
                import shutil
                shutil.copyfileobj(source, output)
    return {'files': results}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--output-dir', type=Path, required=True)
    parser.add_argument('--stem', default='资料整理')
    parser.add_argument('--formats', default='pdf,xlsx,md')
    parser.add_argument('--font')
    args = parser.parse_args()
    try:
        result = export(json.loads(args.input.read_text(encoding='utf-8-sig')), args.output_dir, args.stem,
                        args.formats.split(','), args.font)
        print(json.dumps(result, ensure_ascii=False))
    except Exception as exc:
        print(json.dumps({'error': str(exc)}, ensure_ascii=False), file=sys.stderr)
        raise SystemExit(1)
